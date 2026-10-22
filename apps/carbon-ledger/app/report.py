"""Inventory report export (XLSX and CSV): totals by scope, every entry with its lineage, and the factors used.

Labels are bilingual (Vietnamese / English) because the report is handed to Vietnamese reviewers and auditors.
"""

from __future__ import annotations

import csv
import io
from datetime import datetime, timezone
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

_HEAD = Font(bold=True, color="FFFFFF")
_FILL = PatternFill("solid", fgColor="157A4A")

ENTRY_COLUMNS = [
    ("period", "Kỳ / Period"), ("scope", "Phạm vi / Scope"), ("facility_code", "Cơ sở / Facility"),
    ("source_id", "Nguồn phát thải / Source"), ("filename", "Chứng từ gốc / Source document"),
    ("document_sha256", "SHA-256 chứng từ / Document SHA-256"), ("quantity", "Lượng hoạt động / Activity"),
    ("unit", "Đơn vị / Unit"), ("factor_value", "Hệ số / Factor"), ("factor_unit", "Đơn vị hệ số / Factor unit"),
    ("factor_version", "Phiên bản hệ số / Factor version"), ("factor_status", "Trạng thái hệ số / Factor status"),
    ("co2e_t", "Phát thải (tCO2e) / Emissions (tCO2e)"), ("reviewed_by", "Người kiểm tra / Reviewed by"),
    ("approved_by", "Người phê duyệt / Approved by"), ("entry_id", "Mã bản ghi / Entry id"),
]


def _plain(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        return value.isoformat()
    return value


def rows(entries: list[dict]) -> list[dict]:
    return [{**e, "co2e_t": (e["co2e_kg"] / 1000).quantize(Decimal("0.001")), "entry_id": str(e["id"])} for e in entries]


def xlsx(tenant_name: str, period: str | None, entries: list[dict], periods: list[dict], factors: list[dict]) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Tổng hợp - Summary"
    data = rows(entries)
    scope1 = sum((r["co2e_t"] for r in data if r["scope"] == 1), Decimal(0))
    scope2 = sum((r["co2e_t"] for r in data if r["scope"] == 2), Decimal(0))
    summary.append(["Báo cáo kiểm kê khí nhà kính (Phạm vi 1 và 2) / GHG inventory (Scope 1 and 2)"])
    summary["A1"].font = Font(bold=True, size=14)
    for label, value in (("Doanh nghiệp / Company", tenant_name), ("Kỳ / Period", period or "Tất cả / All"),
                         ("Xuất lúc / Generated", datetime.now(timezone.utc).isoformat(timespec="seconds")),
                         ("Phạm vi 1 (tCO2e) / Scope 1", float(scope1)), ("Phạm vi 2 (tCO2e) / Scope 2", float(scope2)),
                         ("Tổng (tCO2e) / Total", float(scope1 + scope2)), ("Số bản ghi / Entries", len(data))):
        summary.append([label, value])
    summary.append([])
    summary.append(["Kỳ / Period", "Trạng thái / Status", "Niêm phong (lá Trillian) / Seal (Trillian leaf)",
                    "Mã băm kỳ / Period digest", "Khoá bởi / Closed by"])
    for p in periods:
        if period and p["period"] != period:
            continue
        summary.append([p["period"], p["status"], p.get("leaf_index"), p.get("digest"), p.get("closed_by")])
    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 30
    summary.column_dimensions["D"].width = 66

    detail = wb.create_sheet("Chi tiết - Entries")
    detail.append([label for _, label in ENTRY_COLUMNS])
    for r in data:
        detail.append([_plain(r.get(key)) for key, _ in ENTRY_COLUMNS])
    factor_sheet = wb.create_sheet("Hệ số - Factors")
    factor_sheet.append(["Khoá / Key", "Phiên bản / Version", "Giá trị / Value", "Đơn vị / Unit", "Trạng thái / Status", "Nguồn / Source"])
    for f in factors:
        factor_sheet.append([f["factor_key"], f["factor_version"], float(f["factor_value"]), f["factor_unit"], f["factor_status"], f["factor_source"]])

    for sheet in (detail, factor_sheet):
        for cell in sheet[1]:
            cell.font, cell.fill, cell.alignment = _HEAD, _FILL, Alignment(wrap_text=True, vertical="top")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 22
        sheet.freeze_panes = "A2"
    factor_sheet.column_dimensions["F"].width = 90
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def csv_bytes(entries: list[dict]) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow([label for _, label in ENTRY_COLUMNS])
    for r in rows(entries):
        writer.writerow([_plain(r.get(key)) for key, _ in ENTRY_COLUMNS])
    return ("﻿" + out.getvalue()).encode("utf-8")  # BOM so Excel opens Vietnamese text correctly
